# Bench 2: ttk.Treeview con 5000 filas: carga, filtrado por búsqueda y memoria del índice de búsqueda.
import time
import tkinter as tk
import tracemalloc
import unicodedata
from tkinter import ttk

from openpyxl import load_workbook

wb = load_workbook("lista.xlsx", read_only=True)
rows = list(wb.worksheets[0].iter_rows(values_only=True))
wb.close()
header, data = rows[0], rows[1:]


def norm(s):  # casefold + quitar tildes
    return "".join(
        c for c in unicodedata.normalize("NFD", s.casefold()) if unicodedata.category(c) != "Mn"
    )


tracemalloc.start()
t = time.perf_counter()
idx = [norm(" ".join("" if v is None else str(v) for v in r)) for r in data]
print(
    f"indice busqueda: {time.perf_counter() - t:.3f}s, {tracemalloc.get_traced_memory()[0] / 1e6:.2f}MB"
)
tracemalloc.stop()
root = tk.Tk()
root.withdraw()
tv = ttk.Treeview(root, columns=("sel",) + header, show="headings")
t = time.perf_counter()
for i, r in enumerate(data):
    tv.insert("", "end", iid=str(i), values=("☑",) + r)
root.update()
print(f"insert 5000: {time.perf_counter() - t:.3f}s")


def filtro_detach(q):
    q = norm(q)
    t = time.perf_counter()
    match = [str(i) for i, s in enumerate(idx) if q in s]
    tv.detach(*tv.get_children())
    tv.move  # noqa
    for pos, iid in enumerate(match):
        tv.reattach(iid, "", pos)
    root.update()
    return time.perf_counter() - t, len(match)


def filtro_rebuild(q):
    q = norm(q)
    t = time.perf_counter()
    tv.delete(*tv.get_children())
    n = 0
    for i, s in enumerate(idx):
        if q in s:
            tv.insert("", "end", iid=str(i), values=("☑",) + data[i])
            n += 1
    root.update()
    return time.perf_counter() - t, n


for q in ("jose", "user12", "bogota", ""):
    print(f"detach/reattach '{q}': {filtro_detach(q)[0]:.3f}s n={filtro_detach(q)[1]}")
for q in ("jose", "user12", ""):
    dt, n = filtro_rebuild(q)
    print(f"delete/insert '{q}': {dt:.3f}s n={n}")
root.destroy()
