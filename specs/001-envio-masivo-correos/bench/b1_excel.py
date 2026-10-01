# Bench 1: generar xlsx de 5000 filas y medir lectura normal vs read_only.
import gc
import random
import string
import time
import tracemalloc

from openpyxl import Workbook, load_workbook

N = 5000
wb = Workbook(write_only=True)
ws = wb.create_sheet("Destinatarios")
ws.append(["Nombre", "Apellido", "Correo", "Empresa", "Cargo", "Ciudad", "Teléfono", "Notas"])
for i in range(N):
    r = lambda k: "".join(random.choices(string.ascii_letters, k=k))
    ws.append(
        [
            f"José {r(5)}",
            r(8),
            f"user{i}@empresa.com",
            r(12),
            r(10),
            "Bogotá",
            f"+57 300 {i:07d}",
            r(40),
        ]
    )
wb.save("lista.xlsx")


def bench(ro):
    gc.collect()
    tracemalloc.start()
    t = time.perf_counter()
    wb = load_workbook("lista.xlsx", read_only=ro, data_only=True)
    rows = [tuple(c for c in row) for row in wb.worksheets[0].iter_rows(values_only=True)]
    if ro:
        wb.close()
    dt = time.perf_counter() - t
    cur, peak = tracemalloc.get_traced_memory()
    tracemalloc.stop()
    del wb
    return dt, peak / 1e6, cur / 1e6, len(rows)


for ro in (False, True):
    dt, peak, cur, n = bench(ro)
    print(f"read_only={ro}: {dt:.2f}s filas={n} pico={peak:.1f}MB retenido_con_filas={cur:.1f}MB")
