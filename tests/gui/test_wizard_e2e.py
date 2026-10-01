"""Recorrido E2E de la GUI real (quickstart B1 + B2) contra el SMTP de prueba.

Maneja los widgets por código dentro del mainloop; los diálogos modales se reemplazan por
registros para no bloquear. La verificación de dominios usa un resolvedor falso (sin red).
"""

import time

import pytest

from envio_correos import config
from envio_correos.core.domain_check import DomainStatus
from envio_correos.core.journal import Journal
from envio_correos.core.message_builder import SendMode
from envio_correos.devtools.fake_smtp import DEV_PASSWORD, DEV_USER, FakeSmtpServer
from envio_correos.gui import main_window
from envio_correos.gui.presenters import recipients as recipients_presenter
from tests.fixtures.make_fixtures import make_lista

pytestmark = pytest.mark.slow


class AllDomainsOk:
    def check_all(self, domains, progress=None, cancel=None):
        result = dict.fromkeys(domains, DomainStatus.OK)
        if progress:
            progress(len(result), len(result))
        return result


@pytest.fixture
def env(tmp_path, monkeypatch, isolated_app_dirs):
    monkeypatch.setattr(config, "MESSAGES_PER_MINUTE", 6000)  # sin esperas reales
    monkeypatch.setattr(recipients_presenter, "DomainChecker", AllDomainsOk)
    srv = FakeSmtpServer(config.dev_cert_dir()).start()
    monkeypatch.setenv(config.ENV_DEV_SMTP, f"{srv.host}:{srv.port}")
    journal = Journal()
    yield srv, journal, make_lista(tmp_path / "lista_20.xlsx", 20)
    journal.close()
    srv.stop()


def test_recorrido_completo_uno_por_uno(env):
    srv, journal, lista = env
    app = main_window.build(journal)
    dialogs, log = [], {}
    app.confirm = lambda m: (dialogs.append(("confirm", m)), True)[1]
    app.show_info = lambda m: dialogs.append(("info", m))
    app.show_error = lambda m: dialogs.append(("error", m))

    def wait_until(cond, timeout=30):
        end = time.time() + timeout
        while time.time() < end:
            app.update()
            if cond():
                return
            time.sleep(0.02)
        raise TimeoutError

    def script():
        try:
            app.on_next()
            v = app.current_view()
            v.email.insert(0, DEV_USER)
            v.password.insert(0, DEV_PASSWORD)
            v._test()
            wait_until(lambda: not v.presenter.testing)
            log["conexion"] = v.status.cget("text")
            app.on_next()
            v = app.current_view()
            v._load(lista)
            wait_until(
                lambda: (
                    v.table.tree.get_children() and v.notice.cget("text") == "Dominios verificados."
                )
            )
            p = v.presenter
            for i in (0, 1, 2):
                v.table._toggle_iid(str(i))
            v.search.insert(0, "bogota")
            v._apply_search()
            log["visibles"] = len(v.table.tree.get_children())
            v.table._toggle_iid(str(next(i for i in p.visible if p.lst.is_selected(i))))
            v._clear_search()
            log["footer"] = v.footer.cget("text")
            log["filas_tras_limpiar"] = len(v.table.tree.get_children())
            app.on_next()
            v = app.current_view()
            v.subject.insert(0, "Novedades")
            v.body.insert("1.0", "Hola")
            app.on_next()
            app.on_next()  # sin modo: debe avisar
            log["sin_modo"] = dialogs[-1]
            app.current_view()._select(SendMode.ONE_BY_ONE)
            app.on_next()
            app.on_next()  # confirmación (stub) → crea la campaña y arranca
            log["confirmacion"] = dialogs[-1]
            wait_until(lambda: app.wizard.index == 7, 30)
            log["resultado"] = app.current_view().presenter.breakdown()
        except Exception as exc:  # pragma: no cover - se reporta en el assert
            log["error"] = repr(exc)
        finally:
            app.after(50, app.destroy)

    app.after(200, script)
    app.mainloop()

    assert "error" not in log, log
    assert log["conexion"] == f"Conexión exitosa con {DEV_USER}."
    assert log["visibles"] == 4
    assert log["filas_tras_limpiar"] == 20
    assert log["footer"].startswith("Se enviará a 16") and "Excluidos 4" in log["footer"]
    assert log["sin_modo"] == ("info", "Elegí una de las dos formas de envío.")
    assert log["confirmacion"][0] == "confirm" and "16 destinatarios" in log["confirmacion"][1]
    assert log["resultado"]["Enviados"] == 16 and log["resultado"]["Excluidos por vos"] == 4
    assert len(srv.messages) == 16 and {len(m.rcpt_tos) for m in srv.messages} == {1}


class DomainsWithTypo:
    def check_all(self, domains, progress=None, cancel=None):
        return {
            d: DomainStatus.NOT_FOUND if d == "empresa.con" else DomainStatus.OK for d in domains
        }


def test_rechazo_y_dominio_invalido_generan_excel_de_fallidos(env, monkeypatch, tmp_path):
    """Quickstart B3."""
    from openpyxl import Workbook

    from envio_correos.core import excel_reader

    srv, journal, _ = env
    monkeypatch.setattr(recipients_presenter, "DomainChecker", DomainsWithTypo)
    srv.behavior.rcpt_responses["rechazar@ejemplo.com"] = "550 5.1.1 User unknown"
    lista = tmp_path / "b3.xlsx"
    wb = Workbook()
    for row in (
        ("Nombre", "Correo"),
        ("Ana", "ana@x.com"),
        ("Rechazo", "rechazar@ejemplo.com"),
        ("Typo", "juan@empresa.con"),
        ("Luis", "luis@x.com"),
    ):
        wb.active.append(row)
    wb.save(lista)

    app = main_window.build(journal)
    log = {}
    app.confirm = lambda m: True
    app.show_info = app.show_error = lambda m: log.setdefault("dialogs", []).append(m)

    def wait_until(cond, timeout=30):
        end = time.time() + timeout
        while time.time() < end:
            app.update()
            if cond():
                return
            time.sleep(0.02)
        raise TimeoutError

    def script():
        try:
            app.on_next()
            v = app.current_view()
            v.email.insert(0, DEV_USER)
            v.password.insert(0, DEV_PASSWORD)
            v._test()
            wait_until(lambda: not v.presenter.testing)
            app.on_next()
            v = app.current_view()
            v._load(lista)
            wait_until(lambda: v.table.tree.get_children() and "dominio" in v.notice.cget("text"))
            log["footer"] = v.footer.cget("text")
            app.on_next()
            v = app.current_view()
            v.subject.insert(0, "Hola")
            v.body.insert("1.0", "Cuerpo")
            app.on_next()
            app.current_view()._select(SendMode.ONE_BY_ONE)
            app.on_next()
            app.on_next()
            wait_until(lambda: app.wizard.index == 7, 30)
            rv = app.current_view()
            log["breakdown"] = rv.presenter.breakdown()
            log["report"] = rv.presenter.report_path
            log["report_lbl"] = rv.report_lbl.cget("text")
        except Exception as exc:  # pragma: no cover
            log["error"] = repr(exc)
        finally:
            app.after(50, app.destroy)

    app.after(200, script)
    app.mainloop()

    assert "error" not in log, log
    assert log["footer"].startswith("Se enviará a 3")
    assert (log["breakdown"]["Enviados"], log["breakdown"]["Fallidos"]) == (2, 1)
    data = excel_reader.read_sheet(log["report"])
    assert [r[1] for r in data.rows] == ["rechazar@ejemplo.com", "juan@empresa.con"]
    assert data.headers == ("Nombre", "Correo", "Estado", "Motivo", "Código de campaña")
    assert log["report"].name in log["report_lbl"]


def test_corregir_y_reintentar_desde_el_resultado(env, monkeypatch, tmp_path):
    """Quickstart B4 + B5: corregir el typo en el Excel de fallidos y reintentar solo eso;
    una fila agregada con una dirección ya enviada aparece desmarcada."""
    from openpyxl import Workbook, load_workbook

    from envio_correos.core.recipients import ValidationStatus
    from envio_correos.gui import os_open

    srv, journal, _ = env
    monkeypatch.setattr(os_open, "open_path", lambda p: True)  # no abrir Excel de verdad
    srv.behavior.rcpt_responses["rechazar@ejemplo.com"] = "550 5.1.1 User unknown"
    lista = tmp_path / "b4.xlsx"
    wb = Workbook()
    for row in (("Nombre", "Correo"), ("Ana", "ana@x.com"), ("Rechazo", "rechazar@ejemplo.com")):
        wb.active.append(row)
    wb.save(lista)

    app = main_window.build(journal)
    log = {}
    app.confirm = lambda m: True
    app.show_info = app.show_error = lambda m: log.setdefault("dialogs", []).append(m)

    def wait_until(cond, timeout=30):
        end = time.time() + timeout
        while time.time() < end:
            app.update()
            if cond():
                return
            time.sleep(0.02)
        raise TimeoutError

    def send_current(subject=None):
        app.on_next()  # Destinatarios → Mensaje
        v = app.current_view()
        if subject is not None:
            v.subject.insert(0, subject)
            v.body.insert("1.0", "Cuerpo")
        log.setdefault("subjects", []).append(v.subject.get())
        app.on_next()  # → Modo
        if app.current_view().presenter.selected is None:
            app.current_view()._select(SendMode.ONE_BY_ONE)
        app.on_next()  # → Revisión
        app.on_next()  # confirma → Envío
        wait_until(lambda: app.wizard.index == 7, 30)

    def script():
        try:
            app.on_next()
            v = app.current_view()
            v.email.insert(0, DEV_USER)
            v.password.insert(0, DEV_PASSWORD)
            v._test()
            wait_until(lambda: not v.presenter.testing)
            app.on_next()
            v = app.current_view()
            v._load(lista)
            wait_until(
                lambda: (
                    v.table.tree.get_children() and v.notice.cget("text") == "Dominios verificados."
                )
            )
            send_current("Novedades")
            rv = app.current_view()
            report = rv.presenter.report_path
            log["code1"] = rv.presenter.code
            srv.messages.clear()

            # La usuaria corrige el typo y agrega una fila con una dirección ya enviada
            book = load_workbook(report)
            ws = book["Fallidos"]
            ws["B2"] = "corregido@ejemplo.org"
            ws.append(("Ana de nuevo", "ana@x.com", None, None, None))
            book.save(report)

            rv._retry()
            dialog = next(
                w
                for w in app.winfo_children()
                if w.winfo_class() == "CTkToplevel" or w.__class__.__name__ == "RetryWaitDialog"
            )
            dialog._continue()
            wait_until(lambda: app.wizard.index == 2)
            v = app.current_view()
            wait_until(
                lambda: (
                    v.table.tree.get_children() and v.notice.cget("text") == "Dominios verificados."
                )
            )
            lst = v.presenter.lst
            log["retry_statuses"] = [lst.status_of(i) for i in range(len(lst))]
            log["retry_footer"] = v.footer.cget("text")
            send_current()
            rv = app.current_view()
            log["code2"] = rv.presenter.code
            log["breakdown2"] = rv.presenter.breakdown()
        except Exception:  # pragma: no cover
            import traceback

            log["error"] = traceback.format_exc()
        finally:
            app.after(50, app.destroy)

    app.after(200, script)
    app.mainloop()

    assert "error" not in log, log["error"]
    assert log["retry_statuses"] == [ValidationStatus.VALID, ValidationStatus.ALREADY_SENT]
    assert log["retry_footer"].startswith("Se enviará a 1")
    assert log["subjects"] == ["Novedades", "Novedades"]  # mensaje precargado en el reintento
    assert log["code2"] == f"{log['code1']}-R2"
    assert [m.rcpt_tos for m in srv.messages] == [["corregido@ejemplo.org"]]
    assert log["breakdown2"]["Enviados"] == 1 and log["breakdown2"]["Excluidos por vos"] == 1


def test_variables_adjuntos_vista_previa_y_prueba(env, tmp_path):
    """US5: insertar {Nombre} con su botón, adjuntar, previsualizar, enviarme una prueba."""
    from email import message_from_bytes
    from email.policy import default as default_policy

    srv, journal, lista = env
    adj = tmp_path / "folleto.txt"
    adj.write_text("contenido del folleto", encoding="utf-8")
    app = main_window.build(journal)
    log = {}
    app.confirm = lambda m: True
    app.show_info = app.show_error = lambda m: log.setdefault("dialogs", []).append(m)

    def wait_until(cond, timeout=30):
        end = time.time() + timeout
        while time.time() < end:
            app.update()
            if cond():
                return
            time.sleep(0.02)
        raise TimeoutError

    def script():
        try:
            app.on_next()
            v = app.current_view()
            v.email.insert(0, DEV_USER)
            v.password.insert(0, DEV_PASSWORD)
            v._test()
            wait_until(lambda: not v.presenter.testing)
            app.on_next()
            v = app.current_view()
            v._load(lista)
            wait_until(
                lambda: (
                    v.table.tree.get_children() and v.notice.cget("text") == "Dominios verificados."
                )
            )
            v.presenter.mark_visible(False)
            v.table.refresh_rows()
            v.table._toggle_iid("0")
            v.table._toggle_iid("1")
            app.on_next()
            v = app.current_view()
            log["botones"] = [b.cget("text") for b in v.vars_frame.winfo_children()][:3]
            v.subject.insert(0, "Hola ")
            v._set_target(v.subject)
            v._insert("Nombre")
            v.body.insert("1.0", "Te esperamos en ")
            v.body.mark_set("insert", "end-1c")
            v._set_target(v.body)
            v._insert("Ciudad")
            v.presenter.add_files([adj])
            v._render_attachments()
            log["adjuntos"] = v.att_summary.cget("text")
            app.on_next()
            app.current_view()._select(SendMode.ONE_BY_ONE)
            app.on_next()
            rv = app.current_view()
            log["preview1"] = rv.preview.get("1.0", "end")
            rv._move(1)
            log["preview2"] = rv.preview.get("1.0", "end")
            rv._send_test()
            wait_until(lambda: rv.test_lbl.cget("text") != "")
            log["prueba"] = rv.test_lbl.cget("text")
            log["prueba_rcpts"] = [m.rcpt_tos for m in srv.messages]
            srv.messages.clear()
            app.on_next()
            wait_until(lambda: app.wizard.index == 7, 30)
        except Exception:  # pragma: no cover
            import traceback

            log["error"] = traceback.format_exc()
        finally:
            app.after(50, app.destroy)

    app.after(200, script)
    app.mainloop()

    assert "error" not in log, log["error"]
    assert log["botones"] == ["Nombre", "Apellido", "Correo"]
    assert log["adjuntos"].startswith("1 adjunto(s)")
    assert "Asunto: Hola José 0" in log["preview1"] and "Te esperamos en Bogotá" in log["preview1"]
    assert "Adjuntos: folleto.txt" in log["preview1"]
    assert "Asunto: Hola María 1" in log["preview2"]
    assert log["prueba"] == "Prueba enviada a tu casilla."
    assert log["prueba_rcpts"] == [[DEV_USER]]
    assert len(srv.messages) == 2
    parsed = [message_from_bytes(m.data, policy=default_policy) for m in srv.messages]
    assert [p["Subject"] for p in parsed] == ["Hola José 0", "Hola María 1"]
    assert all([a.get_filename() for a in p.iter_attachments()] == ["folleto.txt"] for p in parsed)
    assert parsed[1].get_body(("plain",)).get_content().strip() == "Te esperamos en Medellín"
