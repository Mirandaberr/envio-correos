import datetime as dt

import pytest
from openpyxl import Workbook, load_workbook

from envio_correos.core import excel_reader
from envio_correos.core.excel_reader import ExcelReadError
from envio_correos.texts import es


def test_lista_hojas_y_elige_la_primera_con_datos(fixtures_dir):
    path = fixtures_dir["varias_hojas"]
    assert excel_reader.list_sheets(path) == ["Portada", "Clientes", "Proveedores"]
    assert excel_reader.read_sheet(path).sheet_name == "Clientes"


def test_lee_hoja_elegida(fixtures_dir):
    data = excel_reader.read_sheet(fixtures_dir["varias_hojas"], "Proveedores")
    assert data.headers == ("Razón social", "Contacto")
    assert data.rows == [("ACME", "compras@acme.com")]


def test_filas_como_tuplas_y_numeros_de_fila(fixtures_dir):
    data = excel_reader.read_sheet(fixtures_dir["lista_20"])
    assert len(data.rows) == 20
    assert all(type(r) is tuple and len(r) == 8 for r in data.rows)
    assert list(data.source_row_numbers[:3]) == [2, 3, 4]


def test_casos_borde(fixtures_dir):
    data = excel_reader.read_sheet(fixtures_dir["bordes"])
    assert data.headers == ("Nombre", "Columna B", "Correo electrónico", "Ciudad")
    # la fila vacía intercalada (fila 3) se descarta y la numeración respeta el Excel
    assert 3 not in data.source_row_numbers
    assert data.rows[0] == ("Ana", "x", "ana@empresa.com", "Bogotá")
    assert len(data.rows) == 8


def test_encabezados_repetidos_se_distinguen(tmp_path):
    path = tmp_path / "rep.xlsx"
    wb = Workbook()
    wb.active.append(("Nombre", "Nombre", "Correo"))
    wb.active.append(("A", "B", "a@x.com"))
    wb.save(path)
    assert excel_reader.read_sheet(path).headers == ("Nombre", "Nombre (2)", "Correo")


def test_conserva_tipos_de_celda(tmp_path):
    path = tmp_path / "tipos.xlsx"
    wb = Workbook()
    wb.active.append(("Correo", "Alta", "Monto"))
    wb.active.append(("a@x.com", dt.datetime(2026, 5, 1, 9, 0), 1500.5))
    wb.save(path)
    assert excel_reader.read_sheet(path).rows == [
        ("a@x.com", dt.datetime(2026, 5, 1, 9, 0), 1500.5)
    ]


def test_filas_cortas_se_completan(tmp_path):
    path = tmp_path / "corta.xlsx"
    wb = Workbook()
    wb.active.append(("A", "B", "C"))
    wb.active.append(("x",))
    wb.save(path)
    assert excel_reader.read_sheet(path).rows == [("x", None, None)]


def test_cierra_el_libro_aunque_falle(monkeypatch, fixtures_dir):
    closed = []
    real = load_workbook

    def spy(*a, **k):
        wb = real(*a, **k)
        orig = wb.close
        wb.close = lambda: (closed.append(True), orig())
        return wb

    monkeypatch.setattr(excel_reader, "load_workbook", spy)
    with pytest.raises(KeyError):
        excel_reader.read_sheet(fixtures_dir["lista_20"], "NoExiste")
    assert closed == [True]


@pytest.mark.parametrize(
    ("name", "content", "message"),
    [
        ("viejo.xls", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, es.FILE_FORMAT_XLS),
        ("datos.csv", b"a,b\n1,2\n", es.FILE_FORMAT_CSV),
        ("protegido.xlsx", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, es.FILE_PROTECTED),
        ("roto.xlsx", b"no es un zip", es.FILE_UNREADABLE),
    ],
)
def test_formatos_no_soportados_con_ayuda(tmp_path, name, content, message):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(ExcelReadError) as exc:
        excel_reader.read_sheet(path)
    assert exc.value.user_message == message


def test_archivo_inexistente(tmp_path):
    with pytest.raises(ExcelReadError) as exc:
        excel_reader.read_sheet(tmp_path / "no.xlsx")
    assert exc.value.user_message == es.FILE_NOT_FOUND


def test_archivo_bloqueado(monkeypatch, tmp_path, fixtures_dir):
    def locked(*a, **k):
        raise PermissionError("en uso")

    monkeypatch.setattr(excel_reader, "load_workbook", locked)
    with pytest.raises(ExcelReadError) as exc:
        excel_reader.read_sheet(fixtures_dir["lista_20"])
    assert exc.value.user_message == es.FILE_LOCKED


def test_libro_sin_datos(tmp_path):
    path = tmp_path / "vacio.xlsx"
    Workbook().save(path)
    with pytest.raises(ExcelReadError) as exc:
        excel_reader.read_sheet(path)
    assert exc.value.user_message == es.FILE_NO_DATA


def test_xml_malicioso_se_rechaza_sin_colgarse(tmp_path):
    """Un .xlsx con entidades XML anidadas ("billion laughs") se rechaza gracias a defusedxml."""
    import zipfile

    import openpyxl.xml

    assert openpyxl.xml.DEFUSEDXML
    good = tmp_path / "bueno.xlsx"
    wb = Workbook()
    wb.active.append(("Correo",))
    wb.active.append(("a@x.com",))
    wb.save(good)
    bomb = (
        '<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY lol "lol">'
        '<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
        '<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">]>'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        '<sheetData><row r="1"><c r="A1" t="inlineStr"><is><t>&lol3;</t></is></c></row>'
        "</sheetData></worksheet>"
    )
    evil = tmp_path / "malicioso.xlsx"
    with zipfile.ZipFile(good) as src, zipfile.ZipFile(evil, "w") as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = bomb.encode()
            dst.writestr(item, data)
    with pytest.raises(ExcelReadError) as exc:
        excel_reader.read_sheet(evil)
    assert exc.value.user_message == es.FILE_UNREADABLE
