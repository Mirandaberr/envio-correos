"""Excel de fallidos (contracts/failed-report-xlsx.md, FR-061)."""

import datetime as dt

import pytest
from openpyxl import load_workbook

from envio_correos import config
from envio_correos.core import report
from envio_correos.core.journal import CampaignSpec, Journal, RecipientRow
from envio_correos.core.journal import RecipientState as S
from envio_correos.core.recipients import EXCLUDED_BY_USER
from envio_correos.texts import es

NOW = dt.datetime(2026, 9, 30, 15, 45)
HEADERS = ("Nombre", "Correo", "Alta")


@pytest.fixture
def journal(tmp_path):
    j = Journal(tmp_path / "j.db")
    yield j
    j.close()


def campaign(journal, tmp_path, headers=HEADERS, source_name="clientes.xlsx"):
    spec = CampaignSpec(
        account_email="yo@empresa.com",
        source_path=str(tmp_path / source_name),
        sheet_name="Hoja",
        headers=headers,
        email_col=1,
        subject="s",
        body="b",
        attachments=(),
        mode="ONE_BY_ONE",
    )
    cid = journal.create_campaign(spec)
    alta = dt.datetime(2026, 1, 2, 9, 30)
    extra = (None,) * (len(headers) - 3)
    rows = [
        RecipientRow(2, "ok@x.com", ("Ok", "ok@x.com", alta, *extra)),
        RecipientRow(3, "falla@x.com", ("Falla", " Falla@X.com ", alta, *extra)),
        RecipientRow(
            4,
            "malo",
            ("Malo", "malo", alta, *extra),
            S.SKIPPED,
            "INVALID_SYNTAX",
            es.VALIDATION["INVALID_SYNTAX"],
        ),
        RecipientRow(
            5,
            "excl@x.com",
            ("Excl", "excl@x.com", alta, *extra),
            S.SKIPPED,
            EXCLUDED_BY_USER,
            es.SKIP_EXCLUDED_BY_USER,
        ),
        RecipientRow(
            6,
            "dup@x.com",
            ("Dup", "dup@x.com", alta, *extra),
            S.SKIPPED,
            "DUPLICATE",
            es.VALIDATION["DUPLICATE"],
        ),
        RecipientRow(7, "pend@x.com", ("Pend", "pend@x.com", alta, *extra)),
        RecipientRow(8, "vuelo@x.com", ("Vuelo", "vuelo@x.com", alta, *extra)),
        RecipientRow(9, "stop@x.com", ("Stop", "stop@x.com", alta, *extra)),
    ]
    n = journal.start_attempt(cid, rows)
    journal.mark_result(cid, n, 0, S.SENT)
    journal.mark_result(
        cid, n, 1, S.FAILED, "RECIPIENT_REJECTED", es.reason_text("RECIPIENT_REJECTED"), "550 5.1.1"
    )
    journal.mark_sending(cid, n, 6)
    journal.mark_result(cid, n, 7, S.SKIPPED, "STOPPED_BY_USER", es.SKIP_STOPPED_BY_USER)
    return cid, n


def read(path):
    wb = load_workbook(path)
    data = [tuple(r) for r in wb[config.REPORT_SHEET_NAME].iter_rows(values_only=True)]
    meta = dict(wb[config.REPORT_META_SHEET_NAME].iter_rows(values_only=True))
    return wb, data, meta


def test_filas_incluidas_y_excluidas(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    path = report.write_report(journal, cid, n, now=NOW)
    _, data, _ = read(path)
    names = [r[0] for r in data[1:]]
    assert names == ["Falla", "Malo", "Pend", "Vuelo", "Stop"]  # orden original


def test_columnas_originales_mas_las_de_la_app(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    _, data, _ = read(report.write_report(journal, cid, n, now=NOW))
    assert data[0] == (*HEADERS, config.COL_STATUS, config.COL_REASON, config.COL_CAMPAIGN)
    falla = data[1]
    assert falla[:3] == ("Falla", " Falla@X.com ", dt.datetime(2026, 1, 2, 9, 30))  # sin normalizar
    assert falla[3:] == (es.REPORT_STATUS_FAILED, es.reason_text("RECIPIENT_REJECTED"), cid)


def test_estados_y_motivos(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    _, data, _ = read(report.write_report(journal, cid, n, now=NOW))
    by_name = {r[0]: (r[3], r[4]) for r in data[1:]}
    assert by_name["Malo"] == (es.REPORT_STATUS_INVALID, es.VALIDATION["INVALID_SYNTAX"])
    assert by_name["Pend"] == (es.REPORT_STATUS_PENDING, es.PENDING_PAUSED)
    assert by_name["Vuelo"] == (es.REPORT_STATUS_UNKNOWN, es.UNKNOWN_IF_SENT)
    assert by_name["Stop"] == (es.REPORT_STATUS_PENDING, es.SKIP_STOPPED_BY_USER)


def test_hoja_meta_oculta(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    wb, _, meta = read(report.write_report(journal, cid, n, now=NOW))
    assert wb.worksheets[0].title == config.REPORT_SHEET_NAME
    assert wb[config.REPORT_META_SHEET_NAME].sheet_state == "hidden"
    assert meta["format_version"] == config.REPORT_FORMAT_VERSION
    assert (meta["campaign_id"], meta["attempt"], meta["email_column"]) == (cid, 1, "Correo")
    assert meta["generated_at"].startswith("2026-09-30T15:45")


def test_encabezado_congelado(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    wb, _, _ = read(report.write_report(journal, cid, n, now=NOW))
    assert wb[config.REPORT_SHEET_NAME].freeze_panes == "A2"


def test_nombre_y_ubicacion(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    path = report.write_report(journal, cid, n, now=NOW)
    assert path == tmp_path / "clientes_fallidos_20260930-1545.xlsx"


def test_reintento_no_duplica_columnas_ni_sufijos(journal, tmp_path):
    headers = (*HEADERS, config.COL_STATUS, config.COL_REASON, config.COL_CAMPAIGN)
    cid, n = campaign(journal, tmp_path, headers, "clientes_fallidos_20260930-1200.xlsx")
    path = report.write_report(journal, cid, n, now=NOW)
    _, data, _ = read(path)
    assert data[0] == (*HEADERS, config.COL_STATUS, config.COL_REASON, config.COL_CAMPAIGN)
    assert path.name == "clientes_fallidos_20260930-1545.xlsx"


def test_codigo_de_intento_en_reintentos(journal, tmp_path):
    cid, _ = campaign(journal, tmp_path)
    n2 = journal.start_attempt(
        cid, [RecipientRow(3, "falla@x.com", ("Falla", "falla@x.com", None))]
    )
    journal.mark_result(cid, n2, 0, S.FAILED, "RECIPIENT_REJECTED", "x")
    _, data, meta = read(report.write_report(journal, cid, n2, now=NOW))
    assert data[1][-1] == f"{cid}-R2" and meta["attempt"] == 2


def test_sin_filas_no_genera_archivo(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    for seq in range(8):
        journal.mark_result(cid, n, seq, S.SENT)
    assert report.write_report(journal, cid, n, now=NOW) is None
    assert report.rows_to_fix(journal, cid, n) == []


def test_destino_bloqueado_lanza_error_especifico(journal, tmp_path, monkeypatch):
    cid, n = campaign(journal, tmp_path)

    def locked(*a, **k):
        raise PermissionError("bloqueado")

    monkeypatch.setattr(report, "_write_file", locked)
    with pytest.raises(report.ReportWriteError) as exc:
        report.write_report(journal, cid, n, now=NOW)
    assert exc.value.suggested_path.name == "clientes_fallidos_20260930-1545.xlsx"


def test_carpeta_alternativa(journal, tmp_path):
    cid, n = campaign(journal, tmp_path)
    other = tmp_path / "otra"
    other.mkdir()
    path = report.write_report(journal, cid, n, now=NOW, target=other / "elegido.xlsx")
    assert path == other / "elegido.xlsx" and path.exists()


def test_textos_que_parecen_formula_quedan_como_texto(journal, tmp_path):
    spec = CampaignSpec(
        "yo@empresa.com",
        str(tmp_path / "f.xlsx"),
        "Hoja",
        ("Nombre", "Correo"),
        1,
        "s",
        "b",
        (),
        "ONE_BY_ONE",
    )
    cid = journal.create_campaign(spec)
    n = journal.start_attempt(
        cid, [RecipientRow(2, "a@x.com", ('=HYPERLINK("http://malo","clic")', "a@x.com"))]
    )
    journal.mark_result(cid, n, 0, S.FAILED, "X", "x")
    path = report.write_report(journal, cid, n, now=NOW)
    cell = load_workbook(path)[config.REPORT_SHEET_NAME]["A2"]
    assert cell.data_type == "s" and cell.value == '=HYPERLINK("http://malo","clic")'
