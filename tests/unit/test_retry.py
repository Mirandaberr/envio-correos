"""Lectura del Excel de fallidos como fuente de un reintento (FR-063, FR-064)."""

import datetime as dt
import hashlib

import pytest
from openpyxl import load_workbook

from envio_correos import config
from envio_correos.core import report, retry
from envio_correos.core.journal import CampaignSpec, Journal, RecipientRow
from envio_correos.core.journal import RecipientState as S
from envio_correos.core.message_builder import AttachmentRef, SendMode
from envio_correos.core.recipients import ValidationStatus as V

NOW = dt.datetime(2026, 9, 30, 10, 0)


@pytest.fixture
def journal(tmp_path):
    j = Journal(tmp_path / "j.db")
    yield j
    j.close()


def failed_report(journal, tmp_path, attachments=()):
    spec = CampaignSpec(
        "yo@empresa.com",
        str(tmp_path / "lista.xlsx"),
        "Hoja",
        ("Nombre", "Correo"),
        1,
        "Hola {Nombre}",
        "Cuerpo",
        tuple(a.to_dict() for a in attachments),
        "ONE_BY_ONE",
    )
    cid = journal.create_campaign(spec)
    rows = [
        RecipientRow(i + 2, a, (n, a))
        for i, (n, a) in enumerate(
            [("Ana", "ana@x.com"), ("Juan", "juan@empresa.con"), ("Luis", "luis@x.com")]
        )
    ]
    n = journal.start_attempt(cid, rows)
    journal.mark_result(cid, n, 0, S.SENT)
    journal.mark_result(cid, n, 1, S.FAILED, "RECIPIENT_REJECTED", "Rechazada")
    journal.mark_result(cid, n, 2, S.FAILED, "RECIPIENT_REJECTED", "Rechazada")
    path = report.write_report(journal, cid, n, now=NOW)
    return cid, path


def edit(path, fn):
    wb = load_workbook(path)
    fn(wb)
    wb.save(path)


def test_asocia_la_campana_por_meta(journal, tmp_path):
    cid, path = failed_report(journal, tmp_path)
    r = retry.load_retry(path, journal)
    assert r.campaign_id == cid
    assert r.template.subject == "Hola {Nombre}" and r.mode is SendMode.ONE_BY_ONE
    assert r.next_attempt == 2 and r.code == f"{cid}-R2"
    assert r.recipients.headers[r.recipients.email_col] == "Correo"
    assert len(r.recipients) == 2


def test_sin_meta_usa_la_columna_codigo(journal, tmp_path):
    cid, path = failed_report(journal, tmp_path)
    edit(path, lambda wb: wb.remove(wb[config.REPORT_META_SHEET_NAME]))
    assert retry.load_retry(path, journal).campaign_id == cid


def test_codigo_con_sufijo_de_reintento(journal, tmp_path):
    cid, path = failed_report(journal, tmp_path)

    def change(wb):
        wb.remove(wb[config.REPORT_META_SHEET_NAME])
        ws = wb[config.REPORT_SHEET_NAME]
        for row in ws.iter_rows(min_row=2):
            row[-1].value = f"{cid}-R3"

    edit(path, change)
    assert retry.load_retry(path, journal).campaign_id == cid


def test_sin_campana_conocida_es_campana_nueva(journal, tmp_path):
    _, path = failed_report(journal, tmp_path)

    def change(wb):
        wb.remove(wb[config.REPORT_META_SHEET_NAME])
        ws = wb[config.REPORT_SHEET_NAME]
        ws.delete_cols(ws.max_column)  # borra "Código de campaña"

    edit(path, change)
    r = retry.load_retry(path, journal)
    assert r.campaign_id is None and r.template is None and r.code is None
    assert len(r.recipients) == 2


def test_ya_enviado_desmarcado_pero_marcable(journal, tmp_path):
    _, path = failed_report(journal, tmp_path)

    def add_row(wb):
        wb[config.REPORT_SHEET_NAME].append(("Ana otra vez", "ANA@x.com", None, None, None))

    edit(path, add_row)
    r = retry.load_retry(path, journal)
    lst = r.recipients
    i = next(i for i in range(len(lst)) if lst.address(i) == "ANA@x.com")
    assert lst.status_of(i) is V.ALREADY_SENT and not lst.is_selected(i)
    assert lst.toggle(i) is True


def test_respeta_filas_editadas_agregadas_y_borradas(journal, tmp_path):
    _, path = failed_report(journal, tmp_path)

    def change(wb):
        ws = wb[config.REPORT_SHEET_NAME]
        ws["B2"] = "juan@empresa.com"  # typo corregido
        ws.delete_rows(3)  # luis eliminado
        ws.append(("Nueva", "nueva@x.com", None, None, None))

    edit(path, change)
    lst = retry.load_retry(path, journal).recipients
    assert [lst.address(i) for i in lst.effective_indices()] == ["juan@empresa.com", "nueva@x.com"]


def test_hoja_meta_no_se_ofrece_como_hoja(journal, tmp_path):
    _, path = failed_report(journal, tmp_path)
    assert config.REPORT_META_SHEET_NAME not in retry.load_retry(path, journal).sheets


def test_adjuntos_ausentes_o_modificados(journal, tmp_path):
    ok = tmp_path / "ok.pdf"
    ok.write_bytes(b"ok")
    changed = tmp_path / "cambiado.pdf"
    changed.write_bytes(b"nuevo")
    refs = (
        AttachmentRef(str(ok), 2, hashlib.sha256(b"ok").hexdigest()),
        AttachmentRef(str(changed), 5, hashlib.sha256(b"viejo").hexdigest()),
        AttachmentRef(str(tmp_path / "borrado.pdf"), 1, "x"),
    )
    _, path = failed_report(journal, tmp_path, refs)
    r = retry.load_retry(path, journal)
    assert {(p.ref.path, p.kind) for p in r.attachment_problems} == {
        (str(changed), "changed"),
        (str(tmp_path / "borrado.pdf"), "missing"),
    }


def test_reporte_de_reintento_respeta_columnas_movidas(journal, tmp_path):
    """Si la usuaria reordena columnas en el Excel de fallidos, el nuevo reporte no corre
    los valores (SC-005): usa los encabezados del archivo del intento."""
    cid, path = failed_report(journal, tmp_path)

    def swap(wb):
        ws = wb[config.REPORT_SHEET_NAME]
        for row in ws.iter_rows():
            row[0].value, row[1].value = row[1].value, row[0].value  # Correo primero

    edit(path, swap)
    r = retry.load_retry(path, journal)
    n = journal.start_attempt(cid, r.recipients.to_attempt_rows(), r.recipients.headers)
    journal.mark_result(cid, n, 0, S.FAILED, "RECIPIENT_REJECTED", "x")
    out = report.write_report(journal, cid, n, now=NOW, target=tmp_path / "r2.xlsx")
    wb = load_workbook(out)
    rows = list(wb[config.REPORT_SHEET_NAME].iter_rows(values_only=True))
    assert rows[0][:2] == ("Correo", "Nombre")
    assert rows[1][:2] == ("juan@empresa.con", "Juan")
