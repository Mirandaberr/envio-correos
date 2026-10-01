"""Generador de archivos Excel de prueba.

Se generan en tiempo de test (no se versionan binarios) para que los casos borde
queden documentados como código.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook

HEADERS_8 = ("Nombre", "Apellido", "Correo", "Empresa", "Cargo", "Ciudad", "Teléfono", "Notas")
CITIES = ("Bogotá", "Medellín", "Cali", "Barranquilla", "Cartagena")


def _row(i: int, domain: str = "empresa.com") -> tuple:
    return (
        f"José {i}" if i % 3 == 0 else f"María {i}",
        f"Pérez{i}",
        f"user{i}@{domain}",
        f"Empresa {i % 50}",
        "Analista",
        CITIES[i % len(CITIES)],
        f"+57 300 {i:07d}",
        "nota de prueba",
    )


def make_lista(path: Path, n: int, domain: str = "empresa.com") -> Path:
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("Destinatarios")
    ws.append(HEADERS_8)
    for i in range(n):
        ws.append(_row(i, domain))
    wb.save(path)
    return path


def make_varias_hojas(path: Path) -> Path:
    wb = Workbook()
    vacia = wb.active
    vacia.title = "Portada"
    datos = wb.create_sheet("Clientes")
    datos.append(("Nombre", "Email"))
    datos.append(("Ana", "ana@empresa.com"))
    datos.append(("Luis", "luis@empresa.com"))
    otra = wb.create_sheet("Proveedores")
    otra.append(("Razón social", "Contacto"))
    otra.append(("ACME", "compras@acme.com"))
    wb.save(path)
    return path


def make_bordes(path: Path) -> Path:
    """Casos borde de la spec: filas vacías, columnas sin nombre, celdas combinadas,
    varias direcciones en una celda, duplicados, espacios, vacíos y typo de dominio."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Lista"
    ws.append(("Nombre", None, "Correo electrónico", "Ciudad"))
    ws.append(("Ana", "x", "ana@empresa.com", "Bogotá"))
    ws.append((None, None, None, None))  # fila vacía intercalada
    ws.append(("Ana bis", "x", "  ANA@Empresa.com ", "Cali"))  # duplicado con mayúsculas y espacios
    ws.append(("Varios", "x", "a@x.com; b@x.com", "Cali"))
    ws.append(("Sin correo", "x", None, "Cali"))
    ws.append(("Mal", "x", "no-es-un-correo", "Cali"))
    ws.append(("Typo", "x", "juan@empresa.con", "Cali"))
    ws.append(("Combinada", "x", "comb@empresa.com", "Cali"))
    ws.merge_cells(start_row=9, start_column=4, end_row=10, end_column=4)
    ws.append(("Último", "x", "ultimo@empresa.com", None))
    wb.save(path)
    return path


def make_all(directory: Path) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    return {
        "lista_20": make_lista(directory / "lista_20.xlsx", 20),
        "lista_5000": make_lista(directory / "lista_5000.xlsx", 5000),
        "varias_hojas": make_varias_hojas(directory / "varias_hojas.xlsx"),
        "bordes": make_bordes(directory / "bordes.xlsx"),
    }
