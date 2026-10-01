"""Punto de entrada: logging, mantenimiento local y lanzamiento de la GUI."""

from __future__ import annotations

import logging
import shutil
import sys
import tempfile
import threading
import time
from pathlib import Path

from envio_correos import config
from envio_correos.logging_setup import setup_logging

log = logging.getLogger("envio_correos")


def cleanup_orphan_mei_dirs(now: float | None = None) -> int:
    """PyInstaller onefile no borra su carpeta temporal `_MEI*` si el proceso muere
    (research R8). Se borran las propias con más de MEI_ORPHAN_MAX_AGE_S de antigüedad,
    nunca la del proceso actual."""
    if not config.is_frozen():
        return 0
    current = Path(getattr(sys, "_MEIPASS", "")).resolve()
    now = now or time.time()
    removed = 0
    for d in Path(tempfile.gettempdir()).glob("_MEI*"):
        try:
            if d.resolve() == current or not d.is_dir():
                continue
            if now - d.stat().st_mtime > config.MEI_ORPHAN_MAX_AGE_S:
                shutil.rmtree(d, ignore_errors=True)
                removed += 1
        except OSError:
            continue
    return removed


def _install_excepthooks() -> None:
    def hook(exc_type, exc, tb):
        log.critical("Error no controlado", exc_info=(exc_type, exc, tb))

    sys.excepthook = hook
    threading.excepthook = lambda a: log.critical(
        "Error no controlado en hilo %s",
        a.thread.name if a.thread else "?",
        exc_info=(a.exc_type, a.exc_value, a.exc_traceback),
    )


SMOKE_FLAG = "--smoke"


def smoke_check() -> int:
    """Verifica en el ejecutable empaquetado lo que más suele romperse al empaquetar
    (quickstart P1). Devuelve 0 si todo está bien; los detalles quedan en el log."""
    failures: list[str] = []

    def check(name: str, fn) -> None:
        try:
            fn()
            log.info("smoke OK: %s", name)
        except Exception:
            log.exception("smoke FALLÓ: %s", name)
            failures.append(name)

    def keyring_backend() -> None:
        import keyring
        import keyring.backends.fail

        backend = keyring.get_keyring()
        log.info("Backend de keyring: %s", type(backend).__name__)
        if isinstance(backend, keyring.backends.fail.Keyring):
            raise RuntimeError("keyring sin backend")

    def dns_types() -> None:
        import dns.rdataclass
        import dns.rdatatype
        from dns.rdata import get_rdata_class

        for t in ("MX", "A", "AAAA"):
            get_rdata_class(dns.rdataclass.IN, dns.rdatatype.from_text(t))

    def xml_protection() -> None:
        import openpyxl.xml

        if not openpyxl.xml.DEFUSEDXML:
            raise RuntimeError("openpyxl sin defusedxml: XML malicioso no se bloquearía")

    def data_dir_writable() -> None:
        probe = config.data_dir() / ".smoke"
        probe.write_text("ok", encoding="utf-8")
        probe.unlink()

    def excel_roundtrip() -> None:
        from openpyxl import Workbook, load_workbook

        path = config.data_dir() / ".smoke.xlsx"
        wb = Workbook(write_only=True)
        wb.create_sheet("x").append(["a@x.com"])
        wb.save(path)
        load_workbook(path, read_only=True).close()
        path.unlink()

    def window() -> None:
        import customtkinter as ctk

        root = ctk.CTk()
        root.after(300, root.destroy)
        root.mainloop()

    for name, fn in (
        ("keyring", keyring_backend),
        ("dns", dns_types),
        ("protección XML", xml_protection),
        ("carpeta de datos", data_dir_writable),
        ("excel", excel_roundtrip),
        ("ventana", window),
    ):
        check(name, fn)
    return 1 if failures else 0


def main(argv: list[str] | None = None) -> int:
    argv = sys.argv[1:] if argv is None else argv
    setup_logging()
    if SMOKE_FLAG in argv:
        return smoke_check()
    _install_excepthooks()
    log.info("Inicio de la aplicación (empaquetada=%s)", config.is_frozen())

    from envio_correos.core.journal import Journal

    journal = Journal()
    purged = journal.purge()
    if purged:
        log.info(
            "Historial: %d campañas de más de %d días borradas",
            purged,
            config.JOURNAL_RETENTION_DAYS,
        )
    interrupted = journal.mark_interrupted_running()
    if interrupted:
        log.warning("Se encontraron %d envíos interrumpidos por un cierre inesperado", interrupted)
    cleanup_orphan_mei_dirs()

    from envio_correos.gui.main_window import run

    try:
        return run(journal, argv)
    finally:
        journal.close()
        log.info("Fin de la aplicación")


if __name__ == "__main__":
    sys.exit(main())
