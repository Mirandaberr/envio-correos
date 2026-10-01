"""Lectura de Excel en modo read-only (research R7.1).

`read_only=True` mantiene la memoria casi constante (bench 1: pico 3,8 MB vs 16,7 MB para
5.000 filas) a cambio de: celdas de solo lectura, sin estilos y la obligación de cerrar el
libro explícitamente — se hace en un `finally`.
"""

from __future__ import annotations

import logging
import zipfile
from array import array
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from envio_correos.texts import es

log = logging.getLogger(__name__)

_OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"  # .xls antiguo o .xlsx cifrado


class ExcelReadError(Exception):
    def __init__(self, user_message: str):
        super().__init__(user_message)
        self.user_message = user_message


@dataclass
class SheetData:
    """Datos de una hoja. Las filas son tuplas (no dicts ni objetos) y los números de fila un
    `array('I')` de 4 bytes por entrada: con miles de filas, evita un objeto por celda."""

    sheet_name: str
    headers: tuple[str, ...]
    rows: list[tuple[Any, ...]]
    source_row_numbers: array


def _check_format(path: Path) -> None:
    if not path.exists():
        raise ExcelReadError(es.FILE_NOT_FOUND)
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raise ExcelReadError(es.FILE_FORMAT_CSV)
    if suffix == ".xls":
        raise ExcelReadError(es.FILE_FORMAT_XLS)


def _open(path: Path):
    _check_format(path)
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        raise ExcelReadError(es.FILE_LOCKED) from None
    except ValueError:
        # defusedxml rechaza XML malicioso (entidades anidadas, etc.) con ValueError
        log.warning("Excel rechazado: contenido XML no válido o peligroso")
        raise ExcelReadError(es.FILE_UNREADABLE) from None
    except (zipfile.BadZipFile, InvalidFileException, KeyError, OSError):
        try:
            head = path.read_bytes()[: len(_OLE_SIGNATURE)]
        except OSError:
            head = b""
        msg = es.FILE_PROTECTED if head == _OLE_SIGNATURE else es.FILE_UNREADABLE
        raise ExcelReadError(msg) from None


def list_sheets(path: Path) -> list[str]:
    wb = _open(Path(path))
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _is_empty(row: Iterable[Any]) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def _headers(raw: tuple[Any, ...]) -> tuple[str, ...]:
    """Encabezados vacíos → "Columna C"; repetidos → "Nombre (2)" (las variables deben ser
    únicas)."""
    seen: dict[str, int] = {}
    out = []
    for i, v in enumerate(raw, 1):
        name = (
            str(v).strip()
            if v is not None and str(v).strip()
            else f"Columna {get_column_letter(i)}"
        )
        count = seen.get(name.casefold(), 0) + 1
        seen[name.casefold()] = count
        out.append(name if count == 1 else f"{name} ({count})")
    return tuple(out)


def _trim_width(raw: tuple[Any, ...]) -> int:
    width = len(raw)
    while width and _is_empty(raw[width - 1 : width]):
        width -= 1
    return width


def _read_ws(ws) -> SheetData | None:
    ws.reset_dimensions()  # algunas apps escriben dimensiones incorrectas (doc. openpyxl)
    header_raw: tuple[Any, ...] | None = None
    width = 0
    rows: list[tuple[Any, ...]] = []
    numbers = array("I")
    for excel_row, raw in enumerate(ws.iter_rows(values_only=True), 1):
        if _is_empty(raw):
            continue
        if header_raw is None:
            width = _trim_width(raw)
            header_raw = raw[:width]
            continue
        row = tuple(raw[:width]) + (None,) * (width - len(raw))
        if not _is_empty(row):
            rows.append(row)
            numbers.append(excel_row)
    if header_raw is None or not rows:
        return None
    return SheetData(ws.title, _headers(header_raw), rows, numbers)


def read_sheet(path: Path, sheet_name: str | None = None) -> SheetData:
    """Lee la hoja indicada o, si no se indica, la primera con datos."""
    path = Path(path)
    wb = _open(path)
    try:
        if sheet_name is not None:
            data = _read_ws(wb[sheet_name])
        else:
            data = next((d for ws in wb.worksheets if (d := _read_ws(ws))), None)
    except (ValueError, zipfile.BadZipFile) as exc:
        # En modo read_only las hojas se parsean al recorrerlas: un XML malicioso o roto
        # aparece recién acá. KeyError (hoja inexistente) se deja pasar: es un error de uso.
        log.warning("No se pudo leer la hoja: %s", type(exc).__name__)
        raise ExcelReadError(es.FILE_UNREADABLE) from None
    finally:
        wb.close()
    if data is None:
        raise ExcelReadError(es.FILE_NO_DATA)
    log.info("Excel leído: %d filas, %d columnas", len(data.rows), len(data.headers))
    return data
