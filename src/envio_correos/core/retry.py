"""Reintento a partir del Excel de fallidos editado a mano (FR-063, FR-064,
contracts/failed-report-xlsx.md "Lectura como entrada de reintento")."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from envio_correos import config
from envio_correos.core import attachments, excel_reader
from envio_correos.core.attachments import AttachmentProblem
from envio_correos.core.journal import Journal, attempt_code
from envio_correos.core.message_builder import AttachmentRef, MessageTemplate, SendMode
from envio_correos.core.recipients import RecipientList

log = logging.getLogger(__name__)

_ATTEMPT_SUFFIX_RE = re.compile(r"-R\d+$")


@dataclass
class RetryLoad:
    path: Path
    sheets: list[str]
    recipients: RecipientList
    campaign_id: str | None = None
    template: MessageTemplate | None = None
    mode: SendMode | None = None
    next_attempt: int | None = None
    attachment_problems: list[AttachmentProblem] = field(default_factory=list)

    @property
    def code(self) -> str | None:
        if self.campaign_id is None or self.next_attempt is None:
            return None
        return attempt_code(self.campaign_id, self.next_attempt)


def read_meta(path: Path) -> dict[str, object]:
    """Hoja oculta `_meta` como dict; vacío si no existe o no se puede leer."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - el archivo ya se validó con excel_reader
        return {}
    try:
        if config.REPORT_META_SHEET_NAME not in wb.sheetnames:
            return {}
        ws = wb[config.REPORT_META_SHEET_NAME]
        return {str(r[0]): r[1] for r in ws.iter_rows(values_only=True) if r and r[0] is not None}
    finally:
        wb.close()


def _campaign_from_column(data: excel_reader.SheetData) -> str | None:
    if config.COL_CAMPAIGN not in data.headers:
        return None
    col = data.headers.index(config.COL_CAMPAIGN)
    for row in data.rows:
        value = row[col]
        if value:
            return _ATTEMPT_SUFFIX_RE.sub("", str(value).strip())
    return None


def _resolve_campaign(
    meta: dict[str, object], data: excel_reader.SheetData, journal: Journal
) -> str | None:
    candidates = []
    if meta.get("format_version") == config.REPORT_FORMAT_VERSION and meta.get("campaign_id"):
        candidates.append(str(meta["campaign_id"]))
    if from_col := _campaign_from_column(data):
        candidates.append(from_col)
    return next((c for c in candidates if journal.get_campaign(c) is not None), None)


def load_retry(path: Path, journal: Journal) -> RetryLoad:
    """Lanza excel_reader.ExcelReadError con un mensaje para la usuaria si no se puede leer."""
    path = Path(path)
    sheets = [s for s in excel_reader.list_sheets(path) if s != config.REPORT_META_SHEET_NAME]
    sheet = config.REPORT_SHEET_NAME if config.REPORT_SHEET_NAME in sheets else None
    data = excel_reader.read_sheet(path, sheet)
    meta = read_meta(path)
    cid = _resolve_campaign(meta, data, journal)
    if cid is None:
        log.info("Reintento sin campaña asociada: se trata como campaña nueva")
        return RetryLoad(path, sheets, RecipientList.from_sheet(data, str(path)))

    stored = journal.get_campaign(cid)
    assert stored is not None
    spec = stored.spec
    lst = RecipientList.from_sheet(
        data, str(path), already_sent=journal.already_sent_addresses(cid)
    )
    email_header = str(meta.get("email_column") or spec.headers[spec.email_col])
    if email_header in lst.headers:
        lst.set_email_column(lst.headers.index(email_header))
    refs = tuple(AttachmentRef.from_dict(a) for a in spec.attachments)
    latest = journal.latest_attempt(cid) or 0
    log.info("Reintento de la campaña %s: %d filas", cid, len(lst))
    return RetryLoad(
        path=path,
        sheets=sheets,
        recipients=lst,
        campaign_id=cid,
        template=MessageTemplate(spec.subject, spec.body, refs),
        mode=SendMode(spec.mode),
        next_attempt=latest + 1,
        attachment_problems=attachments.problems(refs),
    )
