"""Excel de fallidos (contracts/failed-report-xlsx.md, FR-061).

Se escribe en modo `write_only` (research R7.2): memoria < 10 MB sin importar el tamaño, a
cambio de que solo se puede agregar filas en orden, la configuración de hoja (freeze panes,
estado oculto) va antes de las filas y el libro se guarda una sola vez.
"""

from __future__ import annotations

import datetime as dt
import io
import logging
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font

from envio_correos import config
from envio_correos.core.journal import Journal, RecipientState, StoredRecipient, attempt_code
from envio_correos.logging_setup import campaign_context
from envio_correos.texts import es

log = logging.getLogger(__name__)

# Motivos de omisión que NO van al reporte: no son errores y no deben volver en un reintento.
NOT_IN_REPORT = frozenset({"EXCLUDED_BY_USER", "DUPLICATE", "ALREADY_SENT"})
INVALID_CODES = frozenset({"INVALID_SYNTAX", "INVALID_DOMAIN", "MULTIPLE_ADDRESSES", "EMPTY"})
PENDING_SKIP_CODES = frozenset({"STOPPED_BY_USER", "INTERRUPTED"})
_PREVIOUS_SUFFIX_RE = re.compile(re.escape(config.REPORT_SUFFIX) + r"\d{8}-\d{4}$")


class ReportWriteError(Exception):
    """No se pudo escribir en la ubicación sugerida (archivo abierto o sin permisos)."""

    def __init__(self, suggested_path: Path):
        super().__init__(str(suggested_path))
        self.suggested_path = suggested_path


@dataclass(frozen=True)
class ReportRow:
    recipient: StoredRecipient
    status: str
    reason: str


def _classify(r: StoredRecipient) -> tuple[str, str] | None:
    if r.state is RecipientState.FAILED:
        return es.REPORT_STATUS_FAILED, r.reason_text or es.reason_text("UNKNOWN")
    if r.state is RecipientState.PENDING:
        return es.REPORT_STATUS_PENDING, es.PENDING_PAUSED
    if r.state is RecipientState.SENDING:
        return es.REPORT_STATUS_UNKNOWN, es.UNKNOWN_IF_SENT
    if r.state is RecipientState.SKIPPED:
        if r.reason_code in INVALID_CODES:
            return es.REPORT_STATUS_INVALID, r.reason_text or ""
        if r.reason_code in PENDING_SKIP_CODES:
            return es.REPORT_STATUS_PENDING, r.reason_text or ""
    return None  # enviado, excluido, duplicado o ya enviado


def rows_to_fix(journal: Journal, campaign_id: str, attempt: int) -> list[ReportRow]:
    out = []
    for r in journal.recipients(campaign_id, attempt):
        c = _classify(r)
        if c is not None:
            out.append(ReportRow(r, *c))
    return out


def default_path(source_path: str, now: dt.datetime) -> Path:
    src = Path(source_path)
    stem = _PREVIOUS_SUFFIX_RE.sub("", src.stem)
    return src.with_name(
        f"{stem}{config.REPORT_SUFFIX}{now.strftime(config.REPORT_TIMESTAMP_FORMAT)}.xlsx"
    )


def write_report(
    journal: Journal,
    campaign_id: str,
    attempt: int,
    *,
    now: dt.datetime | None = None,
    target: Path | None = None,
) -> Path | None:
    """Genera el Excel de fallidos. Devuelve la ruta, o None si no hay filas a corregir.
    Lanza ReportWriteError si no se puede escribir en la ubicación (la GUI ofrece otra)."""
    with campaign_context(attempt_code(campaign_id, attempt)):
        return _write_report(journal, campaign_id, attempt, now=now, target=target)


def _write_report(
    journal: Journal,
    campaign_id: str,
    attempt: int,
    *,
    now: dt.datetime | None,
    target: Path | None,
) -> Path | None:
    rows = rows_to_fix(journal, campaign_id, attempt)
    if not rows:
        return None
    stored = journal.get_campaign(campaign_id)
    assert stored is not None
    spec = stored.spec
    now = now or dt.datetime.now()
    path = target or default_path(spec.source_path, now)
    code = attempt_code(campaign_id, attempt)
    attempt_headers = journal.attempt_headers(campaign_id, attempt)
    # Las columnas que agregó la app en un reporte anterior se reemplazan, no se duplican.
    keep = [i for i, h in enumerate(attempt_headers) if h not in config.APP_ADDED_COLUMNS]
    headers = [attempt_headers[i] for i in keep]

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(config.REPORT_SHEET_NAME)
    ws.freeze_panes = "A2"
    bold = Font(bold=True)
    ws.append([_bold(ws, h, bold) for h in (*headers, *config.REPORT_APP_COLUMNS)])
    for row in rows:
        values: tuple[Any, ...] = row.recipient.values
        ws.append(
            [
                *(_safe(ws, values[i] if i < len(values) else None) for i in keep),
                row.status,
                row.reason,
                code,
            ]
        )
    meta = wb.create_sheet(config.REPORT_META_SHEET_NAME)
    meta.sheet_state = "hidden"
    for item in (
        ("format_version", config.REPORT_FORMAT_VERSION),
        ("campaign_id", campaign_id),
        ("attempt", attempt),
        (
            "email_column",
            spec.headers[spec.email_col] if spec.email_col < len(spec.headers) else "",
        ),
        ("generated_at", now.isoformat(timespec="seconds")),
    ):
        meta.append(item)
    # Se guarda primero en memoria: así openpyxl siempre finaliza y borra sus temporales del
    # modo write_only aunque el destino esté bloqueado (si fallara a mitad, quedarían abiertos).
    # El reporte solo trae las filas a corregir, por lo que el buffer es chico.
    buffer = io.BytesIO()
    wb.save(buffer)
    try:
        _write_file(path, buffer.getvalue())
    except OSError as exc:
        log.warning("No se pudo guardar el Excel de fallidos: %s", type(exc).__name__)
        raise ReportWriteError(path) from None
    journal.set_report_path(campaign_id, attempt, str(path))
    log.info("Excel de fallidos generado con %d filas", len(rows))
    return path


def _safe(ws, value: Any) -> Any:
    """openpyxl guarda como FÓRMULA cualquier texto que empiece con "=". Un texto del Excel
    original (p. ej. `=HYPERLINK(...)`) volvería activo en el reporte: se fuerza como texto."""
    if isinstance(value, str) and value.startswith("="):
        cell = WriteOnlyCell(ws, value=value)
        cell.data_type = "s"
        return cell
    return value


def _write_file(path: Path, data: bytes) -> None:
    path.write_bytes(data)


def _bold(ws, value: str, font: Font) -> WriteOnlyCell:
    cell = WriteOnlyCell(ws, value=value)
    cell.font = font
    return cell
